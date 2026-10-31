"""报表导出实测（模板驱动版）：detail/report/personal/inquiry 四类型。"""
import json
import random
import urllib.error
import urllib.request
from io import BytesIO

from openpyxl import load_workbook

BASE = "http://127.0.0.1:8000"
YEAR, MONTH = 2026, 9
ok = fail = 0
made = []


def req(m, p, b=None, t=None):
    h = {"Content-Type": "application/json"}
    if t:
        h["Authorization"] = "Bearer " + t
    r = urllib.request.Request(BASE + p, data=json.dumps(b).encode() if b else None, headers=h, method=m)
    try:
        with urllib.request.urlopen(r) as x:
            raw = x.read()
            return x.status, (json.loads(raw) if raw and b"{" == raw[:1] else raw)
    except urllib.error.HTTPError as e:
        raw = e.read()
        return e.code, (json.loads(raw) if raw else None)


def ck(l, c, detail=""):
    global ok, fail
    print(("[OK]   " if c else "[FAIL] ") + l + (f"  -> {detail}" if detail else ""))
    ok += c
    fail += not c


def login(phone, pw):
    _, d = req("POST", "/api/auth/login", {"phone": phone, "password": pw})
    return d["access_token"]


def get_wb(url, token):
    r = urllib.request.Request(BASE + url, headers={"Authorization": "Bearer " + token})
    with urllib.request.urlopen(r) as x:
        return load_workbook(BytesIO(x.read()))


AT = login("13800000002", "admin123")

# 清残留 + 造演示成交
_, custs = req("GET", "/api/customers?page=1&page_size=100", t=AT)
for it in custs.get("items", []):
    if "导出测试" in it["company_name"]:
        req("DELETE", f"/api/customers/{it['id']}", t=AT)
_, c1 = req("POST", "/api/customers", {"company_name": "导出测试A", "contact_name": "T", "phone": f"138{random.randint(10000000, 99999999)}"}, AT)
made.append(c1["id"])
_, d1 = req("POST", f"/api/customers/{c1['id']}/deals", {
    "product": "直角封口贴标机", "product_line": "right_angle", "quantity": 2, "unit": "台",
    "unit_price": 4500, "deal_date": f"{YEAR}-{MONTH:02d}-01", "sales_no": "示例号",
}, AT)
ck("① 造演示成交", d1.get("id") is not None)

# ② detail：4 sheet 全对齐
wb = get_wb(f"/api/exports/deals?type=detail&year={YEAR}&month={MONTH}", AT)
ck("② detail 4 sheet", wb.sheetnames == ["销售汇总明细表", "Sheet1", "销售汇总明细表 (2)", "大设备"], str(wb.sheetnames))
ws = wb["销售汇总明细表"]
heads = [ws.cell(1, c).value for c in range(1, 19)]
ck("③ 明细 18 列样板原序", heads[0] == "日期" and heads[1] == "销售手机" and heads[12] == "备注（回扣）" and heads[17] == "一句话总结客户购买原因")
found = any(ws.cell(r, 5).value == "导出测试A" for r in range(2, ws.max_row + 1))
ck("④ 全量 sheet 含演示成交", found)
ws_big = wb["大设备"]
ck("⑤ 大设备 sheet 存在", ws_big.max_row >= 1)

# ③ report：sheet 名 + 目标列
wb2 = get_wb(f"/api/exports/deals?type=report&year={YEAR}&month={MONTH}", AT)
ck("⑥ report sheet 名对齐", wb2.sheetnames == ["2026全年", f"26-{MONTH}", f"26-{MONTH}数据分类"], str(wb2.sheetnames))
ws2 = wb2[f"26-{MONTH}"]
ck("⑦ 26-9 表头（业务人员/台数/业绩）", ws2.cell(2, 1).value == "业务人员" and ws2.cell(2, 2).value == "台数")
# 找目标列（行1 含"总业绩目标"）
tcol = next(c for c in range(1, ws2.max_column + 1) if ws2.cell(1, c).value == "总业绩目标")
ck("⑧ 目标/达成列在位", ws2.cell(1, tcol + 2).value == "达成占比")
ws3 = wb2[f"26-{MONTH}数据分类"]
ck("⑨ 数据分类 A1=销售人员", ws3.cell(1, 1).value == "销售人员")
labels = [ws3.cell(r, 1).value for r in range(2, ws3.max_row + 1)]
ck("⑩ 行标签=销售收入（直角）格式", any(str(x).startswith("销售收入（直角）") for x in labels), str(labels[:3]))

# ⑪ personal 底稿
wb3 = get_wb(f"/api/exports/deals?type=personal&year={YEAR}&month={MONTH}&user_id=1", AT)
sn3 = wb3.sheetnames[0]
ck("⑪ personal 单 sheet", len(wb3.sheetnames) == 1 and "26-9" in sn3, str(wb3.sheetnames))
ck("⑫ personal 行标签含销售收入（直角）", any(ws3cell.value and "销售收入（直角）" in str(ws3cell.value) for row in wb3[sn3].iter_rows(max_col=1) for ws3cell in row))

# ⑬ inquiry 骨架
r = urllib.request.Request(BASE + "/api/exports/deals?type=inquiry", headers={"Authorization": "Bearer " + AT})
with urllib.request.urlopen(r) as x:
    wb4 = load_workbook(BytesIO(x.read()))
ck("⑬ inquiry 骨架 26年9月", "26年9月" in wb4.sheetnames)

# ⑭ 权限：sales 不能导 report
_, st = req("POST", "/api/auth/login", {"phone": "13800000001", "password": "admin123"})
code, _ = req("GET", f"/api/exports/deals?type=report&year={YEAR}&month={MONTH}", t=st["access_token"])
ck("⑭ sales 导 report 403", code == 403)

# 清理
for it in req("GET", "/api/customers?page=1&page_size=100", t=AT)[1]["items"]:
    if "导出测试" in it["company_name"]:
        req("DELETE", f"/api/customers/{it['id']}", t=AT)
print(f"\n===== 通过 {ok} / 失败 {fail} =====（测试数据已清理）")
