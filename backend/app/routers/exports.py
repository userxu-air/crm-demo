"""报表导出（模板驱动）：加载样板清空版模板 → 填 CRM 数据 → 输出。

detail     销售明细表汇总.xlsx   4 sheet（全量/透视/当月/大设备）
report     业绩汇总.xlsx     月 sheet（按天数匹配模板复制）+ 数据分类
personal   个人日报.xlsx         单人月报底稿（B 层：收入/数量自动填，特殊行人工补）
inquiry    询盘统计.xlsx         月份骨架（人工填粉丝数）
"""
import calendar
import io
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import load_workbook
from sqlalchemy.orm import Session, joinedload

from app.config import settings
from app.database import get_db
from app.deps import get_current_user
from app.models import Deal, Role, User
from app.schemas import PRODUCT_LINE_TEXT

from openpyxl.styles import Alignment, Border, Font, Side

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT_YAHEI = Font(name="微软雅黑", size=12)   # 明细数据行（样板实测）
FONT_SONG = Font(name="宋体", size=14)        # 业绩汇总数据格（样板实测）
CENTER = Alignment(horizontal="center", vertical="center")

router = APIRouter(prefix="/api/exports", tags=["exports"])

TPL_DIR = Path(settings.TEMPLATE_DIR)
TEXT2ENUM = {v: k for k, v in PRODUCT_LINE_TEXT.items()}


def _tpl(name: str):
    p = TPL_DIR / name
    if not p.exists():
        raise HTTPException(500, f"模板缺失: {name}")
    return load_workbook(p)


def _resp(wb, filename: str) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


def _ddate(d: Deal) -> date:
    v = d.deal_date
    return v.date() if isinstance(v, datetime) else v


def _month_range(year: int, month: int):
    last = calendar.monthrange(year, month)[1]
    return date(year, month, 1), date(year, month, last)


@router.get("/deals")
def export_deals(
    type: str = "detail",
    year: int = 0,
    month: int = 0,
    user_id: int | None = None,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    now = datetime.now()
    year, month = year or now.year, month or now.month
    if type not in ("detail", "report", "personal", "inquiry"):
        raise HTTPException(400, "type 仅支持 detail/report/personal/inquiry")
    if type != "inquiry" and user.role != Role.admin:
        raise HTTPException(403, "报表仅管理员可导出")
    if type == "personal":
        target_uid = user_id or user.id
        target = db.get(User, target_uid)
        if not target:
            raise HTTPException(404, "目标账号不存在")
        deals = (
            db.query(Deal)
            .options(joinedload(Deal.user), joinedload(Deal.customer))
            .filter(Deal.user_id == target_uid, *_month_range_filters(year, month))
            .order_by(Deal.deal_date, Deal.id)
            .all()
        )
        wb = _personal_wb(deals, target, year, month)
        return _resp(wb, f"个人日报_{target.name}_{year}-{month:02d}.xlsx")
    if type == "inquiry":
        return _resp(_tpl("询盘统计.xlsx"), "询盘统计.xlsx")

    q = db.query(Deal).options(joinedload(Deal.user), joinedload(Deal.customer))
    if user.role == Role.sales:
        q = q.filter(Deal.user_id == user.id)
    if type == "detail":
        all_deals = q.order_by(Deal.deal_date, Deal.id).all()
        start, end = _month_range(year, month)
        month_deals = [d for d in all_deals if start <= _ddate(d) <= end]
        big = [d for d in all_deals if _is_big(d)]
        wb = _detail_fill(all_deals, month_deals, big, year, month)
        return _resp(wb, "销售明细表汇总.xlsx")

    start, end = _month_range(year, month)
    deals = (
        q.filter(Deal.deal_date >= start, Deal.deal_date <= end)
        .order_by(Deal.deal_date, Deal.id)
        .all()
    )
    wb = _report_fill(db, deals, year, month)
    return _resp(wb, "业绩汇总.xlsx")


def _month_range_filters(year: int, month: int):
    start, end = _month_range(year, month)
    return (Deal.deal_date >= start, Deal.deal_date <= end)


def _is_big(d: Deal) -> bool:
    return d.product_line in {"vision", "box_folding", "roller", "precision", "big_right_angle"} or (
        d.unit == "台" and d.unit_price >= 10000
    )


def _fill_detail_sheet(ws, deals, row_height=None):
    """往模板明细 sheet（表头已就位）追加流水行（样式照抄样板：微软雅黑12+边框+居中）。"""
    for d in deals:
        cust = d.customer
        ws.append([
            _ddate(d),
            d.sales_no or "",
            cust.region if cust else "",
            "",
            cust.company_name if cust else "",
            d.product,
            d.unit,
            d.quantity,
            d.unit_price,
            d.amount,
            d.invoice_type,
            d.account,
            d.remark,
            "复购" if d.repurchase else "",
            "",
            PRODUCT_LINE_TEXT.get(d.product_line, ""),
            d.created_at.date() if d.created_at else "",
            "",
        ])
    # 数据行样式照抄样板：微软雅黑 12 + 四边框 + 居中；行高按样板（大设备 33）
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=18):
        for cell in row:
            cell.font = FONT_YAHEI
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
        if row_height:
            ws.row_dimensions[row[0].row].height = row_height
    widths = [11, 9, 8, 8, 24, 20, 6, 7, 9, 10, 10, 14, 18, 6, 26, 11, 11, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.freeze_panes = "A2"


def _detail_fill(all_deals, month_deals, big_deals, year: int, month: int):
    wb = _tpl("销售明细表汇总.xlsx")

    ws1 = wb["销售汇总明细表"]
    _fill_detail_sheet(ws1, all_deals)
    tr = ws1.max_row + 1
    ws1.cell(tr, 1, "合计").font = ws1.cell(1, 1).font.copy(bold=True)

    ws2 = wb["Sheet1"]
    units = sorted({d.unit for d in all_deals})
    # 模板透视区固定 4 行标签（卷/枚/…），按实际单位重写
    ws2.cell(1, 1, "求和项:成交金额")
    ws2.cell(2, 1, "单位")
    r = 3
    for u in units:
        ws2.cell(r, 1, u)
        rep = round(sum(d.amount for d in all_deals if d.unit == u and d.repurchase), 2)
        no = round(sum(d.amount for d in all_deals if d.unit == u and not d.repurchase), 2)
        ws2.cell(r, 2, rep or None)
        ws2.cell(r, 3, no or None)
        ws2.cell(r, 4, round(rep + no, 2))
        r += 1
    ws2.cell(r, 1, "总计")
    ws2.cell(r, 2, round(sum(d.amount for d in all_deals if d.repurchase), 2))
    ws2.cell(r, 3, round(sum(d.amount for d in all_deals if not d.repurchase), 2))
    ws2.cell(r, 4, round(sum(d.amount for d in all_deals), 2))

    ws3 = wb["销售汇总明细表 (2)"]
    _fill_detail_sheet(ws3, month_deals)
    if month_deals:
        tr3 = ws3.max_row + 1
        ws3.cell(tr3, 1, "合计")

    ws4 = wb["大设备"]
    _fill_detail_sheet(ws4, big_deals, row_height=33)
    if big_deals:
        tr4 = ws4.max_row + 1
        ws4.cell(tr4, 1, "合计")
    return wb


# ---------- report：业绩汇总（模板月 sheet 复制填充） ----------

def _tpl_month_sheets(wb, year: int):
    """解析模板月 sheet：{名: (月号, 天数)}。天数按目标年计算（闰年兼容）。"""
    out = {}
    for sn in wb.sheetnames:
        m = re.match(r"^(\d{2})-(\d{1,2})$", sn.strip())
        if m:
            month = int(m.group(2))
            out[sn] = (month, calendar.monthrange(year, month)[1])
    return out


def _report_fill(db: Session, deals, year: int, month: int):
    wb = _tpl("业绩汇总.xlsx")
    tpl_sheets = _tpl_month_sheets(wb, year)
    days = calendar.monthrange(year, month)[1]
    # 选天数 ≥ 目标月天数的模板月 sheet（29 天模板可覆盖 28/29 天）
    cand = [(sn, m, d) for sn, (m, d) in tpl_sheets.items() if d >= days]
    if not cand:
        raise HTTPException(500, f"模板中无 {days} 天月份 sheet")
    src_name = min(cand, key=lambda x: x[2])[0]
    src = wb[src_name]
    days_src = next(d for sn, _, d in cand if sn == src_name)

    src_cls_name = next((s for s in wb.sheetnames if "数据分类" in s), None)

    # 复制后先删源 sheet 再重命名（避免同名冲突被加后缀）
    ws = wb.copy_worksheet(src)
    del wb[src_name]
    ws.title = f"{str(year)[2:]}-{month}"
    if src_cls_name:
        ws_cls = wb.copy_worksheet(wb[src_cls_name])
        del wb[src_cls_name]
        ws_cls.title = f"{str(year)[2:]}-{month}数据分类"
    else:
        ws_cls = None

    # 清模板月 sheet 的旧数据（保留公式/结构），填当月
    # days_src 已在上方解析
    _reset_month_sheet(ws, days_src, days, year, month)

    # 业绩名单：活跃 + 参与业绩汇总开关打开的账号（样板式全员列出）
    row_users = {u.id: u for u in db.query(User).filter(User.is_active == True, User.in_sales_report == True).all()}  # noqa: E712
    targets = {u.id: (u.monthly_target or 0) for u in row_users.values()}

    agg: dict[int, dict[int, tuple[int, float]]] = {}
    for d in deals:
        units = d.quantity if d.unit == "台" else 0
        t0, a0 = agg.setdefault(d.user_id, {}).get(_ddate(d).day, (0, 0.0))
        agg[d.user_id][_ddate(d).day] = (t0 + units, a0 + d.amount)

    total_col = 2 + days_src * 2  # 模板列结构（含多余天列，合计列位置按模板）
    row = 3
    for uid, u in sorted(row_users.items(), key=lambda kv: kv[1].name):
        ws.cell(row, 1, u.name)
        tot_t = tot_a = 0
        for day in range(1, days + 1):
            t, a = agg.get(uid, {}).get(day, (0, 0.0))
            if t:
                cell = ws.cell(row, 1 + day * 2 - 1, t)
                cell.font = FONT_SONG
                cell.border = BORDER
                cell.alignment = CENTER
            if a:
                cell = ws.cell(row, 2 + day * 2 - 1, round(a, 2))
                cell.font = FONT_SONG
                cell.border = BORDER
            tot_t += t
            tot_a += a
        target = targets.get(uid, 0)
        ws.cell(row, total_col, tot_t or None)
        ws.cell(row, total_col + 1, round(tot_a, 2) or None)
        ws.cell(row, total_col + 2, target or None)
        if target > 0:
            remain = round(target - tot_a, 2)
            ws.cell(row, total_col + 3, remain if remain > 0 else None)
            ws.cell(row, total_col + 4, round(tot_a / target, 4))
        row += 1

    ws.cell(row, 1, "合计")
    for day in range(1, days + 1):
        t = sum(agg.get(uid, {}).get(day, (0, 0.0))[0] for uid in row_users)
        a = sum(agg.get(uid, {}).get(day, (0, 0.0))[1] for uid in row_users)
        if t:
            ws.cell(row, 1 + day * 2 - 1, t)
        if a:
            ws.cell(row, 2 + day * 2 - 1, round(a, 2))
    ws.cell(row, total_col, sum(agg.get(uid, {}).get(day, (0, 0.0))[0] for uid in row_users for day in range(1, days + 1)))

    # 数据分类填充：重写表头销售名 + 矩阵值
    if src_cls_name:
        _fill_classify(ws_cls, deals, [(u.id, u.name) for u in row_users.values()])

    # 删除其余模板 sheet，只留 2026全年 + 本月两个
    keep = {"2026全年", ws.title, ws_cls.title if src_cls_name else None}
    for sn in list(wb.sheetnames):
        if sn not in keep:
            del wb[sn]
    return wb


def _reset_month_sheet(ws, days_src: int, days_target: int, year: int, month: int):
    """清模板月 sheet 数据区（含 A 列旧人名，行动态重建）并把日期行改写为目标月。"""
    for row in ws.iter_rows(min_row=3, min_col=1):
        for cell in row:
            if cell.value is not None and not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = None
    ws["A1"] = "日期"
    # 日期行（r1 每日跨 2 列合并，左列写日期值）
    for day in range(1, days_target + 1):
        cell = ws.cell(1, 1 + day * 2 - 1, date(year, month, day))
        cell.number_format = "m月d日"


def _fill_classify(ws, deals, users_list):
    # 重写第 1 行销售名 + 合计列，清旧值与模板残留列
    last_col = len(users_list) + 2
    ws.cell(1, 1, "销售人员")
    for j, (_, name) in enumerate(users_list, 2):
        ws.cell(1, j, name)
    ws.cell(1, last_col, "合计")
    for c in range(last_col + 1, ws.max_column + 1):  # 清模板旧名单列（含残留公式）
        for r in range(1, ws.max_row + 1):
            ws.cell(r, c).value = None
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=last_col):
        for cell in row:
            if cell.value is not None and not (isinstance(cell.value, str) and cell.value.startswith("=")):
                cell.value = None

    agg: dict[tuple, float] = {}
    qty: dict[tuple, float] = {}
    for d in deals:
        key = (d.product_line, d.user_id)
        agg[key] = agg.get(key, 0) + d.amount
        qty[key] = qty.get(key, 0) + d.quantity

    for row in ws.iter_rows(min_row=2, max_col=1):
        label = ws.cell(row[0].row, 1).value
        if not label:
            continue
        m = re.match(r"销售收入（(.+)）", str(label))
        r = row[0].row
        if m:
            enum = TEXT2ENUM.get(m.group(1))
            if not enum:
                continue
            for j, (uid, _) in enumerate(users_list, 2):
                v = agg.get((enum, uid))
                if v:
                    ws.cell(r, j, round(v, 2))
            ws.cell(r, last_col, round(sum(v for (p, _), v in agg.items() if p == enum), 2))
        elif (m2 := re.match(r"销售数量（(.+)）", str(label))):
            enum = TEXT2ENUM.get(m2.group(1))
            if not enum:
                continue
            for j, (uid, _) in enumerate(users_list, 2):
                v = qty.get((enum, uid))
                if v:
                    ws.cell(r, j, v)
            ws.cell(r, last_col, round(sum(v for (p, _), v in qty.items() if p == enum), 2))


# ---------- personal：个人日报底稿 ----------

def _personal_wb(deals, target: User, year: int, month: int):
    wb = _tpl("个人日报.xlsx")
    src = wb["26-9"]
    ws = wb.copy_worksheet(src)
    ws.title = f"{str(year)[2:]}-{month}（{target.name}）"
    del wb["26-9"]

    days = calendar.monthrange(year, month)[1]
    # 日期行重写
    for day in range(1, min(days, 31) + 1):
        c = ws.cell(1, 1 + day, date(year, month, day))
        c.number_format = 'm"月"d"日"'

    # 按天聚合
    by_day: dict[int, list[Deal]] = {}
    for d in deals:
        by_day.setdefault(_ddate(d).day, []).append(d)

    def sum_line(day, kind, line_text):
        enum = TEXT2ENUM.get(line_text)
        total = 0
        for d in by_day.get(day, []):
            if enum and d.product_line == enum:
                total += d.amount if kind == "amount" else d.quantity
            elif enum is None and line_text in (d.product or ""):
                total += d.amount if kind == "amount" else d.quantity
        return round(total, 2) if kind == "amount" else total

    for row in ws.iter_rows(min_row=2, max_col=1):
        label = row[0].value
        if not label:
            continue
        r = row[0].row
        m = re.match(r"(销售收入|销售数量)（(.+)）", str(label))
        if not m:
            continue  # 合计/附属段保留模板公式或留空
        kind = "amount" if m.group(1) == "销售收入" else "qty"
        line_text = m.group(2)
        for day in range(1, days + 1):
            col = 1 + day
            if col > 32:
                break
            v = sum_line(day, kind, line_text)
            if v:
                ws.cell(r, col, v)
    return wb
