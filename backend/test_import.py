# -*- coding: utf-8 -*-
"""一键导入功能实测：预检/正式导入/归属/查重/占位号/权限/模板下载。"""
import json
import random
import urllib.error
import urllib.request
from urllib.parse import quote

BASE = "http://127.0.0.1:8000"
ok = fail = 0


def ck(label, cond, detail=""):
    global ok, fail
    print(("✅" if cond else "❌") + " " + label + (f"  [{detail}]" if detail else ""))
    ok += bool(cond)
    fail += (not cond)


def login(phone, pw):
    r = urllib.request.Request(BASE + "/api/auth/login",
                               data=json.dumps({"phone": phone, "password": pw}).encode(),
                               headers={"Content-Type": "application/json"}, method="POST")
    with urllib.request.urlopen(r) as x:
        return json.loads(x.read())["access_token"]


def upload(path, filepath, token, dry):
    boundary = "----importtest"
    body = (
        f"--{boundary}\r\n"
        'Content-Disposition: form-data; name="file"; filename="'
        + filepath.split("/")[-1] + '"\r\n'
        "Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n"
    ).encode() + open(filepath, "rb").read() + f"\r\n--{boundary}--\r\n".encode()
    h = {"Authorization": "Bearer " + token,
         "Content-Type": f"multipart/form-data; boundary={boundary}"}
    sep = "&" if "?" in path else "?"
    r = urllib.request.Request(BASE + path + (f"{sep}dry_run=true" if dry else ""), data=body, headers=h, method="POST")
    try:
        with urllib.request.urlopen(r, timeout=30) as x:
            return x.status, json.loads(x.read())
    except urllib.error.HTTPError as e:
        return e.code, e.read()[:200]


def req(method, path, body=None, token=None):
    h = {"Content-Type": "application/json"}
    if token:
        h["Authorization"] = "Bearer " + token
    r = urllib.request.Request(BASE + path, data=json.dumps(body).encode() if body else None,
                               headers=h, method=method)
    try:
        with urllib.request.urlopen(r, timeout=20) as x:
            raw = x.read()
            return x.status, (json.loads(raw) if raw[:1] in (b"{", b"[") else raw)
    except urllib.error.HTTPError as e:
        return e.code, e.read()[:200]


AT = login("13800000002", "admin123")
# 清残留（上次中断可能留下测试客户；总量感知翻全页，按「导入」前缀识别含变体）
_total = req("GET", "/api/customers?page=1&page_size=1", token=AT)[1]["total"]
_victims = []
for _pg in range(1, (_total + 99) // 100 + 1):
    _, _l = req("GET", f"/api/customers?page={_pg}&page_size=100", token=AT)
    _victims += [it["id"] for it in _l.get("items", []) if it["company_name"].startswith("导入")]
for _vid in _victims:
    req("DELETE", f"/api/customers/{_vid}", token=AT)
rand = random.randint(10000000, 99999999)

# ── 成交流水 ──
code, r = upload("/api/import/deals-flow", "test_import_deals.xlsx", AT, dry=True)
ck("流水预检（不写库）：3 行/2 新客户/3 成交", code == 200 and r["total_rows"] == 3
   and r["customers_created"] == 2 and r["deals_created"] == 3, f"deals={r.get('deals_created')}")
ck("  └ 未知销售手机告警", any("13800000000" in w for w in r.get("warnings", [])))
code, r2 = upload("/api/import/deals-flow", "test_import_deals.xlsx", AT, dry=False)
ck("流水正式导入", code == 200 and r2["deals_created"] == 3 and not r2["dry_run"])

code, lst = req("GET", "/api/customers?page=1&page_size=50&keyword=" + quote("导入测试"), token=AT)
ck("  └ 客户落库（A厂+B贸易）", lst["total"] == 2)
a = next((i for i in lst["items"] if i["company_name"] == "导入测试A厂"), None)
b = next((i for i in lst["items"] if i["company_name"] == "导入测试B贸易"), None)
ck("  └ A厂归属 admin（销售=13800000002 是 admin）", a and a["owner_name"])
ck("  └ B贸易电话留空（界面显示'-'）", b and not b["phone"], str(b["phone"]) if b else "")
code, flow = req("GET", "/api/deals?page=1&page_size=50&keyword=" + quote("导入测试"), token=AT)
imported = [i for i in flow["items"] if "导入测试" in i["company_name"]]
ck("  └ 成交 3 笔落库（9000/6000空算/800 复购）",
   len(imported) == 3 and {i["amount"] for i in imported} == {9000, 6000, 800})
line_hit = next((i for i in imported if i["company_name"] == "导入测试A厂" and i["amount"] == 800), None)
ck("  └ 产品简称→产品线枚举（标签耗材）", line_hit and line_hit["product_line"] == "consumable")

# 二次导入同名客户 → 查重不新建
code, r3 = upload("/api/import/deals-flow", "test_import_deals.xlsx", AT, dry=False)
ck("重复导入 → 客户按名查重不再新建", code == 200 and r3["customers_created"] == 0, f"created={r3['customers_created']} matched={r3['customers_matched']}")

# ── 商机表 ──
code, r4 = upload("/api/import/opportunities", "test_import_opps.xlsx", AT, dry=True)
ck("商机预检：2 行/2 新客户/1 成交", code == 200 and r4["total_rows"] == 2 and r4["deals_created"] == 1, str(r4)[:120])
code, r5 = upload("/api/import/opportunities", "test_import_opps.xlsx", AT, dry=False)
ck("商机正式导入", code == 200 and r5["customers_created"] == 2)

code, lst2 = req("GET", "/api/customers?page=1&page_size=50&keyword=" + quote("导入商机"), token=AT)
jia = next((i for i in lst2["items"] if i["company_name"] == "导入商机甲公司"), None)
yi = next((i for i in lst2["items"] if i["company_name"] == "导入商机乙工作室"), None)
ck("  └ 已成交商机 → won + 成交记录", jia and jia["stage"] == "won")
ck("  └ 未成交商机 → following + 默认 3 天回访", yi and yi["stage"] == "following" and yi.get("next_followup_at"))
code, flow2 = req("GET", "/api/deals?page=1&page_size=20&keyword=" + quote("导入商机甲公司"), token=AT)
jia_deal = next((i for i in flow2["items"] if i["company_name"] == "导入商机甲公司"), None)
ck("  └ 商机成交记录=报价 35000（视觉线）", jia_deal and jia_deal["amount"] == 35000 and jia_deal["product_line"] == "vision")

# ── 权限与模板 ──
code, _ = req("GET", "/api/import/template/deals-flow", token=AT)
ck("模板下载（成交流水）", code == 200)
code, _ = req("GET", "/api/import/template/opportunities", token=AT)
ck("模板下载（商机）", code == 200)
_, d = req("POST", "/api/auth/login", {"phone": "13800000001", "password": "admin123"})
ST = d["access_token"]
code, _ = req("GET", "/api/import/template/deals-flow", token=ST)
ck("sales 下载模板（分账号设计允许）", code == 200)
r = urllib.request.Request(BASE + "/api/import/deals-flow", data=b"x", method="POST")
try:
    urllib.request.urlopen(r)
    ck("未登录导入 → 拒绝", False)
except urllib.error.HTTPError as e:
    ck("未登录导入 → 401/422", e.code in (401, 422))

# 清理
code, lst3 = req("GET", "/api/customers?page=1&page_size=100", token=AT)
n = 0
for it in lst3["items"]:
    if it["company_name"].startswith("导入测试") or it["company_name"].startswith("导入商机"):
        req("DELETE", f"/api/customers/{it['id']}", token=AT)
        n += 1
print(f"\n清理测试客户 {n} 个")
print(f"===== 一键导入实测：{ok} 通过 / {fail} 失败 =====")

# ── 分账号导入：sales 导自己的 ──
# 建销售账号
code, u = req("POST", "/api/users", {"name": "分账号销售", "phone": f"135{random.randint(10000000,99999999)}", "password": "s123456", "role": "sales"}, AT)
SUID, SPH = u["id"], u["phone"]
# 造一个销售自己账号的流水文件（销售手机列=销售自己的号 + 一个别人的号混入）
from openpyxl import Workbook
from datetime import datetime
wb = Workbook(); ws = wb.active; ws.title = "销售汇总明细表"
ws.append(["日期","销售手机","地区","行业","客户名","产品","单位","数量","单价","成交金额","开票类型","收款账户","备注（回扣）","复购","发货地址","产品简称","添加日期","一句话总结客户购买原因"])
ws.append([datetime(2026,9,22), SPH, "广东佛山", "", "分账号测试客户", "圆瓶贴标机", "台", 1, 5000, 5000, "", "", "", "", "", "圆瓶", "", ""])
ws.append([datetime(2026,9,22), "13800000000", "", "", "分账号混入客户", "平面贴标机", "台", 1, 4000, 4000, "", "", "", "", "", "平面", "", ""])
wb.save("test_import_self.xlsx")
STK = login(SPH, "s123456")
code, r = upload("/api/import/deals-flow", "test_import_self.xlsx", STK, dry=False)
ck("sales 自导入成功（不再 403）", code == 200 and r["deals_created"] == 2, str(r)[:100])
code, lst = req("GET", "/api/customers?page=1&page_size=100&keyword=" + quote("分账号"), token=STK)
ck("  └ sales 只看到自己导入的客户（2 条全归本人）", lst["total"] == 2 and all(i["owner_name"] == "分账号销售" for i in lst["items"]))
# 混入别人手机号的行：sales 导入时强制归本人（不看列），无告警
code, r2 = upload("/api/import/deals-flow", "test_import_self.xlsx", STK, dry=True)
ck("  └ sales 重导预检无「归属管理员」告警（列被忽略）", not any("归属管理员" in w for w in r2.get("warnings", [])))
# admin 指定归属账号导入
code, r3 = upload("/api/import/deals-flow?owner_id=%d" % SUID, "test_import_self.xlsx", AT, dry=True)
ck("admin 指定 owner_id 预检成功", code == 200)
# boss 拒绝
code, d = req("POST", "/api/auth/login", {"phone": "13800000003", "password": "boss123"})
code, _ = req("GET", "/api/import/template/deals-flow", token=d["access_token"])
ck("boss 下载模板 → 403", code == 403)
# 清理
code, lst = req("GET", "/api/customers?page=1&page_size=100&keyword=" + quote("分账号"), token=AT)
for it in lst["items"]:
    req("DELETE", f"/api/customers/{it['id']}", token=AT)
req("DELETE", f"/api/users/{SUID}", token=AT)
import os
os.remove("test_import_self.xlsx")
print("分账号场景清理完成")
print(f"===== 追加：分账号导入 =====")

# ── 归一化匹配 + 溯源批次 ──
# 造文件：A厂 名称用全角+空格变体再导一次 → 应命中已有客户而非新建
wb3 = Workbook(); ws3 = wb3.active; ws3.title = "销售汇总明细表"
ws3.append(["日期","销售手机","地区","行业","客户名","产品","单位","数量","单价","成交金额","开票类型","收款账户","备注（回扣）","复购","发货地址","产品简称","添加日期","一句话总结客户购买原因"])
ws3.append([datetime(2026,9,25), "13800000002", "广东东莞", "包装", "　导入 测试Ａ 厂　", "配件", "个", 1, 100, 100, "", "", "", "", "", "其他", "", ""])
wb3.save("test_import_norm.xlsx")
_q1 = req("GET", "/api/customers?page=1&page_size=100&keyword=" + quote("导入测试A厂"), token=AT)
print("[探针] A厂 在库:", [(i["id"], i["company_name"]) for i in _q1.get("items", [])])
code, r = upload("/api/import/deals-flow", "test_import_norm.xlsx", AT, dry=False)
print("[探针] norm 响应:", json.dumps({k: r[k] for k in ("customers_created", "customers_matched", "deals_created", "duplicates")}, ensure_ascii=False))
ck("名称归一匹配（全角+空格变体 → 命中已有 A厂，不新建）", code == 200 and r["customers_created"] == 0, f"created={r['customers_created']} matched={r['customers_matched']}")
ck("  └ 返回批次号（溯源）", bool(r.get("batch")) and len(r["batch"]) == 32)
# 疑似重复提示
wb4 = Workbook(); ws4 = wb4.active; ws4.title = "销售汇总明细表"
ws4.append(["日期","销售手机","地区","行业","客户名","产品","单位","数量","单价","成交金额","开票类型","收款账户","备注（回扣）","复购","发货地址","产品简称","添加日期","一句话总结客户购买原因"])
ws4.append([datetime(2026,9,26), "13800000002", "广东东莞", "包装", "导入测试A厂（客户经理）", "配件", "个", 1, 100, 100, "", "", "", "", "", "其他", "", ""])
wb4.save("test_import_similar.xlsx")
code, r2 = upload("/api/import/deals-flow", "test_import_similar.xlsx", AT, dry=True)
ck("疑似重复提示（含已有名）", any("疑似重复" in w for w in r2.get("warnings", [])), str(r2.get("warnings", []))[:80])
# 成交重复明细列出（重导 norm 文件 → 1 行判重且有明细）
code, r3 = upload("/api/import/deals-flow", "test_import_norm.xlsx", AT, dry=True)
ck("成交重复跳过有明细（客户｜日期｜产品｜金额）", r3.get("duplicates", 0) >= 1 and any("成交重复跳过" in e for e in r3.get("errors", [])))
import os
for f in ("test_import_norm.xlsx", "test_import_similar.xlsx"):
    if os.path.exists(f): os.remove(f)
# 清理（总量感知翻全页，按「导入」前缀识别含变体）
_t = req("GET", "/api/customers?page=1&page_size=1", token=AT)[1]["total"]
_v = []
for _pg in range(1, (_t + 99) // 100 + 1):
    _, _l = req("GET", f"/api/customers?page={_pg}&page_size=100", token=AT)
    _v += [it["id"] for it in _l.get("items", []) if it["company_name"].startswith("导入")]
for _vid in _v:
    req("DELETE", f"/api/customers/{_vid}", token=AT)
print("清理完成", len(_v))
